"""
backend/tests/test_cumulative_merger.py — Tests for Cumulative Excel Append Engine
"""
import shutil
import pytest
import openpyxl
from backend.config import settings
from backend.excel.cumulative_merger import (
    CumulativeMerger,
    StudentRosterRecord,
    DateMarkItem,
)


@pytest.fixture
def clean_subject():
    subj_id = "test_ss_divb_cumulative"
    subj_dir = settings.SUBJECTS_DIR / subj_id
    if subj_dir.exists():
        shutil.rmtree(subj_dir, ignore_errors=True)
    subj_dir.mkdir(parents=True, exist_ok=True)
    yield subj_id
    shutil.rmtree(subj_dir, ignore_errors=True)


@pytest.fixture
def sample_roster():
    return [
        StudentRosterRecord(sr_no=1, roll_no="24108B0001", name="VEDANT PATOLE", batch=1),
        StudentRosterRecord(sr_no=2, roll_no="24108B0002", name="ARYA SURYAVANSHI", batch=1),
        StudentRosterRecord(sr_no=3, roll_no="24108B0005", name="AYUSH SAWANT", batch=1),
    ]


def test_first_scan_initializes_master_excel(clean_subject, sample_roster):
    merger = CumulativeMerger(clean_subject)

    records = [
        DateMarkItem(roll_no="24108B0001", date="2026-09-09", status="P"),
        DateMarkItem(roll_no="24108B0002", date="2026-09-09", status="P"),
        DateMarkItem(roll_no="24108B0005", date="2026-09-09", status="A"),
    ]

    res = merger.merge(
        roster=sample_roster,
        new_dates=["2026-09-09"],
        records=records,
        subject_metadata={"code": "SS", "name": "System Software"},
    )

    assert res.excel_path.exists()
    assert res.added_dates == ["2026-09-09"]

    wb = openpyxl.load_workbook(res.excel_path, data_only=False)
    ws = wb.active

    # Check header
    assert ws["E4"].value == "9/9/26"
    assert ws["F4"].value == "TOTAL"
    assert ws["G4"].value == "HELD"
    assert ws["H4"].value == "ATT %"

    # Check marks
    assert ws["E5"].value == "P"
    assert ws["E6"].value == "P"
    assert ws["E7"].value == "A"

    # Check formulas
    assert '=COUNTIF(E5:E5,"P")' in ws["F5"].value
    assert '=COUNTIF(E5:E5,"P")+COUNTIF(E5:E5,"A")' in ws["G5"].value
    assert '=IF(G5=0,"",F5/G5*100)' in ws["H5"].value

    wb.close()


def test_second_scan_appends_chronologically(clean_subject, sample_roster):
    merger = CumulativeMerger(clean_subject)

    # 1. Week 1: 2026-09-09
    merger.merge(
        roster=sample_roster,
        new_dates=["2026-09-09"],
        records=[
            DateMarkItem(roll_no="24108B0001", date="2026-09-09", status="P"),
            DateMarkItem(roll_no="24108B0002", date="2026-09-09", status="P"),
            DateMarkItem(roll_no="24108B0005", date="2026-09-09", status="A"),
        ],
    )

    # 2. Week 2: 2026-09-23
    res2 = merger.merge(
        roster=sample_roster,
        new_dates=["2026-09-23"],
        records=[
            DateMarkItem(roll_no="24108B0001", date="2026-09-23", status="P"),
            DateMarkItem(roll_no="24108B0002", date="2026-09-23", status="P"),
            DateMarkItem(roll_no="24108B0005", date="2026-09-23", status="P"),
        ],
    )

    assert merger.backup_path.exists()  # Backup created
    assert res2.added_dates == ["2026-09-23"]

    wb = openpyxl.load_workbook(res2.excel_path, data_only=False)
    ws = wb.active

    # Check both date columns exist
    assert ws["E4"].value == "9/9/26"
    assert ws["F4"].value == "23/9/26"
    assert ws["G4"].value == "TOTAL"
    assert ws["H4"].value == "HELD"
    assert ws["I4"].value == "ATT %"

    # Verify formula updated to span E..F
    assert '=COUNTIF(E5:F5,"P")' in ws["G5"].value
    assert '=COUNTIF(E5:F5,"P")+COUNTIF(E5:F5,"A")' in ws["H5"].value

    wb.close()


def test_out_of_order_date_sorted_chronologically(clean_subject, sample_roster):
    merger = CumulativeMerger(clean_subject)

    # First add Sept 23
    merger.merge(
        roster=sample_roster,
        new_dates=["2026-09-23"],
        records=[DateMarkItem(roll_no="24108B0001", date="2026-09-23", status="P")],
    )

    # Then insert earlier date Sept 9
    res = merger.merge(
        roster=sample_roster,
        new_dates=["2026-09-09"],
        records=[DateMarkItem(roll_no="24108B0001", date="2026-09-09", status="P")],
    )

    wb = openpyxl.load_workbook(res.excel_path, data_only=False)
    ws = wb.active

    # 9/9/26 must be in column E, 23/9/26 in column F
    assert ws["E4"].value == "9/9/26"
    assert ws["F4"].value == "23/9/26"
    wb.close()


def test_idempotent_duplicate_merge(clean_subject, sample_roster):
    merger = CumulativeMerger(clean_subject)

    records = [
        DateMarkItem(roll_no="24108B0001", date="2026-09-09", status="P"),
        DateMarkItem(roll_no="24108B0002", date="2026-09-09", status="A"),
    ]

    # Merge once
    merger.merge(roster=sample_roster, new_dates=["2026-09-09"], records=records)

    # Merge identical again
    res = merger.merge(roster=sample_roster, new_dates=["2026-09-09"], records=records)

    assert len(res.added_dates) == 0  # No new dates added
    assert len(res.conflicts) == 0    # No conflicts
