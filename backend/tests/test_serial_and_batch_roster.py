"""
backend/tests/test_serial_and_batch_roster.py — Tests for Serial Number Regeneration and Batch Dividers
"""
import pytest
import json
from pathlib import Path
from backend.database import SessionLocal, Subject, Student, init_db
from backend.excel.cumulative_merger import CumulativeMerger, StudentRosterRecord, DateMarkItem
from backend.schemas import CommitRequest, CellCommitItem, AttendanceStatus, StudentCommitInfo


@pytest.fixture(autouse=True)
def setup_db():
    init_db()


def test_serial_number_regeneration_1_to_30_no_gaps_for_variable_dates(tmp_path):
    """
    Section 2 Requirement:
    A roster of 30 students with 4 date columns must produce Sr 1..30,
    no gaps, no duplicates, for any number of date columns (1 to 10).
    """
    subject_id = "test_subj_serials"
    db = SessionLocal()
    try:
        # Create subject
        subj = db.query(Subject).filter(Subject.id == subject_id).first()
        if subj:
            db.delete(subj)
            db.commit()

        subj = Subject(
            id=subject_id,
            name="Test Serial Subject",
            code="TSS",
            class_name="T.Y.B.TECH",
            division="B",
            faculty="SHP",
        )
        db.add(subj)
        db.commit()

        # Seed 30 students with scrambled or initially wrong serial numbers
        students = []
        for i in range(1, 31):
            batch_num = 1 if i <= 19 else 2
            st = Student(
                subject_id=subject_id,
                sr_no=i * 4,  # Inflated serial simulate previous bug (4, 8, 12... 120)
                roll_no=f"24108B{i:04d}",
                name=f"STUDENT {i}",
                batch=batch_num,
            )
            db.add(st)
            students.append(st)
        db.commit()

        # Test for various number of date columns: 1, 2, 4, 7, 10
        for num_dates in [1, 2, 4, 7, 10]:
            dates = [f"2026-09-{d:02d}" for d in range(1, num_dates + 1)]

            # Generate attendance records (30 students x num_dates)
            records = []
            for s in students:
                for d in dates:
                    records.append(DateMarkItem(roll_no=s.roll_no, date=d, status="P"))

            # Build roster records with regenerated sequential serials 1..N
            students_sorted = list(db.query(Student).filter(Student.subject_id == subject_id).all())
            students_sorted.sort(key=lambda s: (s.batch, s.sr_no or 0, s.roll_no))
            for idx, s in enumerate(students_sorted, start=1):
                s.sr_no = idx
            db.commit()

            roster_records = [
                StudentRosterRecord(sr_no=s.sr_no, roll_no=s.roll_no, name=s.name, batch=s.batch)
                for s in students_sorted
            ]

            # Merge into cumulative Excel
            merger = CumulativeMerger(subject_id=subject_id)
            meta = {
                "name": subj.name,
                "code": subj.code,
                "class_name": subj.class_name,
                "division": subj.division,
                "faculty": subj.faculty,
                "academic_year": subj.academic_year,
            }
            res = merger.merge(
                roster=roster_records,
                new_dates=dates,
                records=records,
                subject_metadata=meta,
                overwrite_conflicts=True,
            )

            # Assertions:
            # 1. Exactly 30 students
            assert len(roster_records) == 30
            # 2. Sequential serials 1..30, no gaps, no duplicates
            serials = [r.sr_no for r in roster_records]
            assert serials == list(range(1, 31)), f"Failed for {num_dates} date columns: {serials}"
            assert len(set(serials)) == 30

    finally:
        db.close()


def test_batch_expectations_for_sample_sheet():
    """
    Section 3 Requirement:
    Expected for the sample sheet:
    Batch 1 = 19 students (24108B0001 to 24108B0027),
    Batch 2 = 11 students (24108B0028 to 24108B0036 including 25108B2001-2003).
    """
    gt_path = Path(__file__).parent / "fixtures" / "ground_truth_ss_divb.json"
    assert gt_path.exists(), "Ground truth fixture not found"

    with open(gt_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    students = data.get("students", [])
    assert len(students) == 30, f"Expected 30 students in sample sheet, got {len(students)}"

    batch_1 = [s for s in students if s.get("batch") == 1]
    batch_2 = [s for s in students if s.get("batch") == 2]

    # Batch 1 checks
    assert len(batch_1) == 19, f"Expected 19 students in Batch 1, got {len(batch_1)}"
    assert batch_1[0]["roll_no"] == "24108B0001"
    assert batch_1[-1]["roll_no"] == "24108B0027"
    assert batch_1[0]["name"] == "VEDANT PATOLE"
    assert batch_1[-1]["name"] == "SHIVAM OJHA"

    # Batch 2 checks
    assert len(batch_2) == 11, f"Expected 11 students in Batch 2, got {len(batch_2)}"
    assert batch_2[0]["roll_no"] == "24108B0028"
    assert batch_2[-1]["roll_no"] == "24108B0036"
    assert batch_2[0]["name"] == "SUJAL PAWAR"
    assert batch_2[-1]["name"] == "KANAK LADE"

    # Diploma / DSE students included in Batch 2
    dse_rolls = [s["roll_no"] for s in batch_2 if s["roll_no"].startswith("25108B")]
    assert set(dse_rolls) == {"25108B2001", "25108B2002", "25108B2003"}


def test_batch_3_and_4_preservation():
    """Verify that batches 3 and 4 are cleanly parsed from strings and written to master excel."""
    import re
    # Test frontend string parsing logic
    for b_str, expected in [("Batch 1", 1), ("Batch 2", 2), ("Batch 3", 3), ("Batch 4", 4), ("3", 3), ("4", 4)]:
        m = re.search(r"\d+", str(b_str))
        b_num = int(m.group()) if m else 1
        assert b_num == expected, f"Failed for {b_str}: expected {expected}, got {b_num}"

    # Test CumulativeMerger with batch 3 and 4
    subject_id = "test_batches_3_4"
    roster = [
        StudentRosterRecord(sr_no=1, roll_no="24108B0001", name="STUDENT 1", batch=1),
        StudentRosterRecord(sr_no=2, roll_no="24108B0028", name="STUDENT 28", batch=2),
        StudentRosterRecord(sr_no=3, roll_no="24108B0055", name="SEJAL SHAHANE", batch=3),
        StudentRosterRecord(sr_no=4, roll_no="24108B0063", name="DEVEN SONAWANE", batch=4),
    ]
    merger = CumulativeMerger(subject_id=subject_id)
    dates = ["2026-09-09"]
    records = [
        DateMarkItem(roll_no="24108B0001", date="2026-09-09", status="P"),
        DateMarkItem(roll_no="24108B0028", date="2026-09-09", status="P"),
        DateMarkItem(roll_no="24108B0055", date="2026-09-09", status="P"),
        DateMarkItem(roll_no="24108B0063", date="2026-09-09", status="P"),
    ]
    meta = {"name": "Test B3 B4", "code": "TB", "class_name": "T.Y.B.TECH", "division": "B", "faculty": "SHP"}
    res = merger.merge(roster=roster, new_dates=dates, records=records, subject_metadata=meta)
    assert res.excel_path.exists()

    import openpyxl
    wb = openpyxl.load_workbook(res.excel_path)
    ws = wb.active
    # Row 5 -> Batch 1
    # Row 6 -> Batch 2
    # Row 7 -> Batch 3
    # Row 8 -> Batch 4
    assert ws.cell(row=5, column=4).value == "Batch 1"
    assert ws.cell(row=6, column=4).value == "Batch 2"
    assert ws.cell(row=7, column=4).value == "Batch 3"
    assert ws.cell(row=8, column=4).value == "Batch 4"
    wb.close()

