"""
backend/tests/test_api_attendance.py — Integration Tests for Scan, Commit & Download Workflow
"""
import io
from pathlib import Path
import openpyxl
import pytest

from backend.config import settings

SAMPLE_IMAGE_PATH = Path("backend/tests/fixtures/sample_sheet_vit.jpg")


import shutil

@pytest.fixture
def test_subject_with_roster(client, ground_truth_data):
    """Seed a test subject with the full 30-student VIT roster."""
    subj_id = f"test_ss_divb_api_{id(client)}"
    subj_dir = settings.SUBJECTS_DIR / subj_id
    if subj_dir.exists():
        shutil.rmtree(subj_dir, ignore_errors=True)

    students_in = [
        {
            "sr_no": s["sr_no"],
            "roll_no": s["roll_no"],
            "name": s["name"],
            "batch": s["batch"],
        }
        for s in ground_truth_data["students"]
    ]

    payload = {
        "id": subj_id,
        "name": "System Software",
        "code": "SS",
        "class_name": "T.Y.B.TECH",
        "division": "B",
        "type": "Theory",
        "faculty": "SHP",
        "academic_year": "2026-27 (Odd)",
        "roster": students_in,
    }
    res = client.post("/api/subjects", json=payload)
    assert res.status_code == 201

    yield subj_id

    if subj_dir.exists():
        shutil.rmtree(subj_dir, ignore_errors=True)


def test_scan_preview_endpoint(client, test_subject_with_roster):
    """Test image upload and AI preview grid generation (mock mode)."""
    with open(SAMPLE_IMAGE_PATH, "rb") as f:
        img_bytes = f.read()

    files = {"file": ("sheet.jpg", io.BytesIO(img_bytes), "image/jpeg")}
    res = client.post(f"/api/subjects/{test_subject_with_roster}/scan?mock=true", files=files)
    assert res.status_code == 200
    preview = res.json()

    assert preview["subject_id"] == test_subject_with_roster
    assert len(preview["date_columns"]) == 4
    assert len(preview["rows"]) == 30
    assert preview["scan_id"] is not None

    # Check first row
    r1 = preview["rows"][0]
    assert r1["roll_no"] == "24108B0001"
    assert len(r1["cells"]) == 4


def test_commit_blocking_rule_on_uncertain(client, test_subject_with_roster):
    """Test that commit is blocked if any cell status is UNCERTAIN."""
    commit_payload = {
        "date_columns": ["2026-09-09"],
        "records": [
            {"roll_no": "24108B0001", "date": "2026-09-09", "status": "P"},
            {"roll_no": "24108B0014", "date": "2026-09-09", "status": "UNCERTAIN"},  # Blocking cell
        ],
    }
    res = client.post(f"/api/subjects/{test_subject_with_roster}/commit", json=commit_payload)
    assert res.status_code == 400
    assert "UNCERTAIN status" in res.json()["detail"]


def test_successful_commit_and_download(client, test_subject_with_roster):
    """Test teacher confirmation committing to SQLite and generating master Excel."""
    # 1. Commit resolved records
    commit_payload = {
        "date_columns": ["2026-09-09"],
        "records": [
            {"roll_no": "24108B0001", "date": "2026-09-09", "status": "P"},
            {"roll_no": "24108B0002", "date": "2026-09-09", "status": "P"},
            {"roll_no": "24108B0005", "date": "2026-09-09", "status": "A"},
        ],
    }
    commit_res = client.post(f"/api/subjects/{test_subject_with_roster}/commit", json=commit_payload)
    assert commit_res.status_code == 200
    commit_data = commit_res.json()
    assert commit_data["added_dates"] == ["2026-09-09"]
    assert "download" in commit_data["master_download_url"]

    # 2. Download master Excel
    down_res = client.get(f"/api/subjects/{test_subject_with_roster}/download")
    assert down_res.status_code == 200
    assert "spreadsheet" in down_res.headers.get("content-type", "")

    # Validate downloaded Excel binary
    wb = openpyxl.load_workbook(io.BytesIO(down_res.content), data_only=False)
    ws = wb.active
    assert ws["E4"].value == "9/9/26"
    assert ws["E5"].value == "P"
    assert ws["E7"].value == "A"
    wb.close()


def test_cumulative_second_commit_appends(client, test_subject_with_roster):
    """Test that scanning/committing a second date appends to the same master Excel."""
    # 1. Commit Week 1
    client.post(f"/api/subjects/{test_subject_with_roster}/commit", json={
        "date_columns": ["2026-09-09"],
        "records": [
            {"roll_no": "24108B0001", "date": "2026-09-09", "status": "P"},
        ],
    })

    # 2. Commit Week 2
    res2 = client.post(f"/api/subjects/{test_subject_with_roster}/commit", json={
        "date_columns": ["2026-09-23"],
        "records": [
            {"roll_no": "24108B0001", "date": "2026-09-23", "status": "P"},
        ],
    })
    assert res2.status_code == 200
    assert res2.json()["added_dates"] == ["2026-09-23"]

    # 3. Download and verify both dates exist
    down_res = client.get(f"/api/subjects/{test_subject_with_roster}/download")
    wb = openpyxl.load_workbook(io.BytesIO(down_res.content), data_only=False)
    ws = wb.active
    assert ws["E4"].value == "9/9/26"
    assert ws["F4"].value == "23/9/26"
    wb.close()
