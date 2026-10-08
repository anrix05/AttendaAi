"""
attendance_writer.py
─────────────────────
Generates a proper VIT-style attendance Excel sheet that
mirrors the structure of the uploaded image.

Sheet layout:
  Row 1 : Title / subject info
  Row 2 : Column headers (Sr No | Roll No | Name | Batch | Date1 | Date2 | ...)
  Row 3+: Student data rows with color-coded attendance marks
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────────────────────
# STYLES
# ─────────────────────────────────────────────────────────────

THIN  = Side(style="thin")
THICK = Side(style="medium")

THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_BORDER = Border(left=THIN, right=THIN, top=THICK, bottom=THICK)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

# Fill colors
FILL_TITLE    = PatternFill("solid", fgColor="1F3864")   # dark blue
FILL_HEADER   = PatternFill("solid", fgColor="2E75B6")   # blue
FILL_SUBHEAD  = PatternFill("solid", fgColor="D6E4F0")   # light blue
FILL_PRESENT  = PatternFill("solid", fgColor="C6EFCE")   # green
FILL_ABSENT   = PatternFill("solid", fgColor="FFC7CE")   # red
FILL_UNCERTAIN = PatternFill("solid", fgColor="FFEB9C")  # yellow
FILL_NM       = PatternFill("solid", fgColor="F2F2F2")   # light gray
FILL_BATCH    = PatternFill("solid", fgColor="E2EFDA")   # light green (batch header)

# Fonts
FONT_TITLE   = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONT_HEADER  = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_PRESENT = Font(name="Calibri", size=10, bold=True, color="375623")
FONT_ABSENT  = Font(name="Calibri", size=10, bold=True, color="9C0006")
FONT_UNCERT  = Font(name="Calibri", size=10, bold=True, color="7D4508")
FONT_NORMAL  = Font(name="Calibri", size=10)
FONT_NM      = Font(name="Calibri", size=10, color="999999")


# ─────────────────────────────────────────────────────────────
# HELPER: apply style to a cell
# ─────────────────────────────────────────────────────────────

def _style(
    ws_cell,
    value,
    fill=None,
    font=None,
    alignment=CENTER,
    border=THIN_BORDER,
):
    ws_cell.value     = value
    if fill:
        ws_cell.fill  = fill
    if font:
        ws_cell.font  = font
    ws_cell.alignment = alignment
    ws_cell.border    = border


# ─────────────────────────────────────────────────────────────
# MAIN WRITER
# ─────────────────────────────────────────────────────────────

def write_attendance_excel(result: dict, output_path: str | Path) -> Path:
    """
    Generate a complete attendance Excel from the pipeline result.

    Parameters
    ----------
    result : dict returned by process_sheet.process()
        Keys: sheet_info, students, decisions
    output_path : where to save the .xlsx file

    Returns
    -------
    Path to saved file.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    sheet_info = result.get("sheet_info") or {}
    students   = result.get("students")   or []
    decisions  = result.get("decisions")  or []

    dates      = sheet_info.get("dates")  or []
    subject    = sheet_info.get("subject") or "Attendance Sheet"
    class_name = sheet_info.get("class")   or ""
    acad_year  = sheet_info.get("academic_year") or ""

    # Build decision lookup: (student_idx, date_idx) → status
    decision_map = {}
    for dec in decisions:
        key = (dec.get("student_idx", -1), dec.get("date_idx", -1))
        decision_map[key] = dec

    # ── Workbook ─────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # ── Column structure ──────────────────────────────────────
    FIXED_COLS = 4   # Sr No | Roll No | Name | Batch
    total_cols = FIXED_COLS + len(dates)

    # ── Row 1: Title ─────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_text = subject
    if class_name:
        title_text += f"  |  {class_name}"
    if acad_year:
        title_text += f"  |  {acad_year}"

    _style(ws.cell(1, 1), title_text, fill=FILL_TITLE, font=FONT_TITLE)
    ws.row_dimensions[1].height = 30

    # ── Row 2: Column headers ─────────────────────────────────
    headers = ["Sr No", "Roll No", "Name", "Batch"] + [
        d if d else f"Date {i+1}" for i, d in enumerate(dates)
    ]

    for col_idx, header in enumerate(headers, start=1):
        _style(
            ws.cell(2, col_idx),
            header,
            fill=FILL_HEADER,
            font=FONT_HEADER,
            border=HEADER_BORDER,
        )

    ws.row_dimensions[2].height = 20

    # ── Rows 3+: Student data ─────────────────────────────────
    current_batch = None
    data_row = 3

    for s_idx, student in enumerate(students):
        # Insert batch separator row when batch changes
        batch_val = str(student.get("batch") or "").strip()
        if batch_val and batch_val != current_batch:
            current_batch = batch_val
            ws.merge_cells(
                start_row=data_row,
                start_column=1,
                end_row=data_row,
                end_column=total_cols,
            )
            _style(
                ws.cell(data_row, 1),
                f"BATCH {batch_val}",
                fill=FILL_BATCH,
                font=Font(name="Calibri", size=10, bold=True),
                border=THIN_BORDER,
            )
            ws.row_dimensions[data_row].height = 16
            data_row += 1

        # Fixed columns
        ws.row_dimensions[data_row].height = 22
        _style(ws.cell(data_row, 1), student.get("sr_no") or s_idx + 1,
               font=FONT_NORMAL, alignment=CENTER, border=THIN_BORDER)
        _style(ws.cell(data_row, 2), student.get("roll_no") or "",
               font=FONT_NORMAL, alignment=CENTER, border=THIN_BORDER)
        _style(ws.cell(data_row, 3), student.get("name") or "",
               font=FONT_NORMAL, alignment=LEFT, border=THIN_BORDER)
        _style(ws.cell(data_row, 4), batch_val or "",
               font=FONT_NORMAL, alignment=CENTER, border=THIN_BORDER)

        # Attendance columns
        for d_idx in range(len(dates)):
            dec = decision_map.get((s_idx, d_idx))

            if dec:
                status     = dec.get("status", "NOT_MARKED")
                confidence = dec.get("confidence", 0.0)
            else:
                # Fallback to Gemini raw if no decision
                att_raw = student.get("attendance") or {}
                g_val   = att_raw.get(str(d_idx + 1), "NM")
                status  = {
                    "P": "PRESENT", "A": "ABSENT",
                    "NM": "NOT_MARKED", "?": "UNCERTAIN"
                }.get(g_val, "NOT_MARKED")
                confidence = 0.90

            cell_col = FIXED_COLS + d_idx + 1
            ws_cell  = ws.cell(data_row, cell_col)

            if status == "PRESENT":
                display = "P"
                _style(ws_cell, display, fill=FILL_PRESENT,
                       font=FONT_PRESENT, border=THIN_BORDER)
            elif status == "ABSENT":
                display = "A"
                _style(ws_cell, display, fill=FILL_ABSENT,
                       font=FONT_ABSENT, border=THIN_BORDER)
            elif status == "UNCERTAIN":
                display = "?"
                _style(ws_cell, display, fill=FILL_UNCERTAIN,
                       font=FONT_UNCERT, border=THIN_BORDER)
            else:
                display = ""
                _style(ws_cell, display, fill=FILL_NM,
                       font=FONT_NM, border=THIN_BORDER)

            # Add confidence as cell comment (optional tooltip)
            if confidence < 0.85 and status != "NOT_MARKED":
                ws_cell.comment = None  # openpyxl comments need extra setup; skip for now

        data_row += 1

    # ── Summary row ───────────────────────────────────────────
    _add_summary_row(ws, data_row, students, dates, decision_map, FIXED_COLS, total_cols)

    # ── Column widths ─────────────────────────────────────────
    ws.column_dimensions["A"].width = 6    # Sr No
    ws.column_dimensions["B"].width = 14   # Roll No
    ws.column_dimensions["C"].width = 24   # Name
    ws.column_dimensions["D"].width = 7    # Batch

    for d_idx in range(len(dates)):
        col_letter = get_column_letter(FIXED_COLS + d_idx + 1)
        ws.column_dimensions[col_letter].width = 8

    # ── Freeze header rows ────────────────────────────────────
    ws.freeze_panes = "E3"   # freeze rows 1-2 and columns A-D

    # ── Page setup ────────────────────────────────────────────
    ws.page_setup.orientation     = "landscape"
    ws.page_setup.fitToWidth      = 1
    ws.page_setup.fitToHeight     = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(output_path)
    print()
    print(f"✅ Excel saved: {output_path}")
    return output_path


# ─────────────────────────────────────────────────────────────
# SUMMARY ROW
# ─────────────────────────────────────────────────────────────

def _add_summary_row(
    ws, row: int, students, dates, decision_map, fixed_cols, total_cols
):
    """Add a TOTAL row at the bottom showing present counts per date."""
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=fixed_cols
    )
    _style(
        ws.cell(row, 1),
        "TOTAL PRESENT",
        fill=FILL_HEADER,
        font=FONT_HEADER,
        border=THIN_BORDER,
    )
    ws.row_dimensions[row].height = 18

    for d_idx in range(len(dates)):
        count = 0
        for s_idx in range(len(students)):
            dec = decision_map.get((s_idx, d_idx))
            if dec and dec.get("status") == "PRESENT":
                count += 1
            elif not dec:
                # fallback
                att = students[s_idx].get("attendance") or {}
                if att.get(str(d_idx + 1)) == "P":
                    count += 1

        col = fixed_cols + d_idx + 1
        _style(
            ws.cell(row, col),
            count,
            fill=FILL_HEADER,
            font=FONT_HEADER,
            border=THIN_BORDER,
        )
