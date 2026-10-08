from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension


# ============================================================
# STYLES
# ============================================================

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="EDE9FE",
)

PRESENT_FILL = PatternFill(
    fill_type="solid",
    fgColor="DCFCE7",
)

ABSENT_FILL = PatternFill(
    fill_type="solid",
    fgColor="FEE2E2",
)

UNCERTAIN_FILL = PatternFill(
    fill_type="solid",
    fgColor="FEF3C7",
)

UNKNOWN_FILL = PatternFill(
    fill_type="solid",
    fgColor="F9FAFB",
)


# ============================================================
# CELL VALUE
# ============================================================

def cell_value(cell):
    text = str(
        cell.get("text", "")
    ).strip()

    status = str(
        cell.get("status", "UNKNOWN")
    ).upper()

    # Attendance status gets priority.
    if status == "PRESENT":
        return "P"

    if status == "ABSENT":
        return "A"

    if status == "UNCERTAIN":
        return "?"

    # OCR text only for cells where there is no
    # attendance mark.
    return text


# ============================================================
# CELL FILL
# ============================================================

def apply_cell_style(excel_cell, data):
    status = str(
        data.get("status", "UNKNOWN")
    ).upper()

    excel_cell.border = THIN_BORDER

    excel_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    if status == "PRESENT":
        excel_cell.fill = PRESENT_FILL
        excel_cell.font = Font(
            bold=True,
        )

    elif status == "ABSENT":
        excel_cell.fill = ABSENT_FILL
        excel_cell.font = Font(
            bold=True,
        )

    elif status == "UNCERTAIN":
        excel_cell.fill = UNCERTAIN_FILL
        excel_cell.font = Font(
            bold=True,
        )

    else:
        excel_cell.fill = UNKNOWN_FILL


# ============================================================
# GENERATE EXCEL
# ============================================================

def generate_excel(result, output_path):
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Attendance"

    matrix = result.get(
        "attendance",
        [],
    )

    if not isinstance(matrix, list):
        matrix = []

    max_columns = 0

    for row in matrix:
        if isinstance(row, list):
            max_columns = max(
                max_columns,
                len(row),
            )

    # --------------------------------------------------------
    # Write matrix
    # --------------------------------------------------------

    for row_index, row in enumerate(
        matrix,
        start=1,
    ):

        if not isinstance(row, list):
            continue

        worksheet.row_dimensions[
            row_index
        ].height = 28

        for column_index, data in enumerate(
            row,
            start=1,
        ):

            if not isinstance(data, dict):
                data = {
                    "text": str(data),
                    "status": "UNKNOWN",
                }

            excel_cell = worksheet.cell(
                row=row_index,
                column=column_index,
            )

            excel_cell.value = cell_value(
                data
            )

            apply_cell_style(
                excel_cell,
                data,
            )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    for column_index in range(
        1,
        max_columns + 1,
    ):

        letter = get_column_letter(
            column_index
        )

        # Wider cells for text/name columns,
        # but still dynamic.
        maximum_length = 12

        if matrix:
            for row in matrix:
                if (
                    isinstance(row, list)
                    and
                    column_index <= len(row)
                ):
                    data = row[
                        column_index - 1
                    ]

                    if isinstance(data, dict):
                        text = str(
                            data.get(
                                "text",
                                "",
                            )
                        )

                        maximum_length = max(
                            maximum_length,
                            min(
                                len(text) + 2,
                                28,
                            ),
                        )

        worksheet.column_dimensions[
            letter
        ].width = maximum_length

    # --------------------------------------------------------
    # Freeze panes
    # --------------------------------------------------------

    worksheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # Print settings
    # --------------------------------------------------------

    worksheet.sheet_view.showGridLines = True

    worksheet.page_setup.orientation = (
        "landscape"
    )

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.print_options.horizontalCentered = True

    worksheet.sheet_properties.outlinePr.summaryBelow = True

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    workbook.save(
        output_path
    )

    return output_path